"""Editable, scoped XLSX report; formulas are authored only in trusted cells."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

BLUE = '3157D5'
CREAM = 'F4F1EB'

def render_report_excel(data, report, guardians):
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    ws = wb.active
    ws.title = 'Report'

    def text(sheet, row, column, value):
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            # Names/comments must never become spreadsheet formulas.
            cell.data_type = 's'
        return cell

    def table(name, headers, rows):
        sheet = wb.create_sheet(name)
        for row, values in enumerate([headers, *rows], 1):
            for col, value in enumerate(values, 1):
                text(sheet, row, col, value)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor=BLUE)
            cell.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
        return sheet

    student = data['student']
    ind = data['indicators']
    primary = sorted((g for g in guardians if g.get('active', 1)),
                     key=lambda g: (not g.get('primary', False), g.get('id', 0)))
    table('Student Lookup', ['Student ID', 'Student name', 'Class', 'Academic year'],
          [[student['id'], student['name'], student['class_name'], student.get('academic_year', '')]])
    table('Guardian Lookup', ['Student ID', 'Guardian name', 'Email', 'Phone', 'Relationship', 'Primary', 'Consent'],
          [[student['id'], g['name'], g['email'], g.get('phone', ''), g.get('relationship', ''),
            bool(g.get('primary')), bool(g.get('consent'))] for g in primary])
    table('Academic Records', ['Subject', 'Score', 'Previous score', 'Change (points)', 'Attendance'],
          [[s['subject'], s.get('score'), s.get('previous'), s.get('delta'), s.get('attendance')]
          for s in data.get('subjects', [])])
    table('Support', ['Session', 'Date', 'Time', 'Room', 'State', 'Attendance', 'Outcome'],
          [[e.get('class', {}).get('topic', ''), e.get('class', {}).get('date', ''),
            e.get('class', {}).get('start_time', ''), e.get('class', {}).get('room', ''),
            e.get('class', {}).get('state', ''), e.get('attendance', ''), e.get('outcome', '')]
           for e in data.get('support', [])])
    readme = table('Read Me', ['Topic', 'Explanation'], [
        ['Purpose', 'Editable MIS sample for one permitted student. Synthetic demonstration data.'],
        ['Academic source', 'Saved report snapshot; editing Excel does not update the website.'],
        ['Guardian source', 'Active guardian contacts at download time; primary contact is first.'],
        ['Lookup', 'Report formulas use exact-match VLOOKUP on Student ID. Excel recalculates when opened.'],
        ['Privacy', 'No passwords, password hashes, session tokens or login secrets are exported.'],
        ['Dates and units', 'Report timestamp is UTC. Scores/attendance are 0–100; change is percentage points.'],
        ['Live data', 'This is an offline workbook, not a live database connection.'],
        ['Review', 'Teacher comments and follow-up can be edited locally. Review before sharing.']])
    readme.column_dimensions['B'].width = 85
    for row in readme.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        readme.row_dimensions[row[0].row].height = 38

    labels = {
        1: 'TEAM BLACKCATS — Student Progress Report', 3: 'Report ID', 4: 'Student ID',
        5: 'Student name', 6: 'Class', 7: 'Academic year', 8: 'Term', 9: 'Report state',
        10: 'Version', 11: 'Generated (UTC)', 13: 'Average score (0–100)',
        14: 'Attendance (0–100)', 15: 'Change (points)', 16: 'Academic status',
        17: 'Review signals', 19: 'Primary guardian', 20: 'Guardian email', 21: 'Guardian phone',
        23: 'Teacher comments — editable', 26: 'Follow-up plan — editable',
        29: 'Note: local edits do not update the saved website report.'}
    for row, label in labels.items():
        text(ws, row, 1, label)
        ws.cell(row, 1).font = Font(bold=True, color=BLUE)
    values = {3: report.id, 4: student['id'], 8: data['term']['name'],
              9: 'Released' if report.released else 'Draft', 10: report.version,
              11: report.generated_at, 13: ind.get('current_average'),
              14: ind.get('attendance'), 15: ind.get('trend_delta'),
              16: ind.get('status'), 17: ', '.join(ind.get('reason_codes', [])),
              23: report.comments, 26: report.follow_up}
    for row, value in values.items():
        text(ws, row, 2, value)
    for row, column in [(5, 2), (6, 3), (7, 4)]:
        ws.cell(row, 2, f'=IFERROR(VLOOKUP($B$4,\'Student Lookup\'!$A$2:$D$2,{column},FALSE),"Not found")')
    last = max(2, len(primary) + 1)
    for row, column in [(19, 2), (20, 3), (21, 4)]:
        ws.cell(row, 2, f'=IFERROR(VLOOKUP($B$4,\'Guardian Lookup\'!$A$2:$G${last},{column},FALSE),"No contact")')
    ws.merge_cells('A1:B1')
    ws['A1'].font = Font(size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=BLUE)
    ws.row_dimensions[1].height = 34
    ws.column_dimensions['A'].width = 37
    ws.column_dimensions['B'].width = 85
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=CREAM)
        ws.row_dimensions[row[0].row].height = 27
    ws.row_dimensions[23].height = 105
    ws.row_dimensions[26].height = 105
    ws.freeze_panes = 'B4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = 'A1:B29'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
